"""
Load-test helper for content_freshness_httpV1.py apply mode (standalone).

  download : save live HTML for every URL in the approved Excel -> <dir>/<slug>.html
  mutate   : copy <src> -> <dst>, appending a unique marker paragraph inside the
             main content of N% of pages, so each mutated page gets a new content_hash.

Typical flow (after a clean freshness run):
  python load_test_fixtures.py download --excel approved.xlsx --out fixtures_base
  python load_test_fixtures.py mutate --src fixtures_base --dst fixtures_mutated --percent 100
  python content_freshness_httpV1.py --excel approved.xlsx --fixture-dir fixtures_mutated        # report
  python content_freshness_httpV1.py --excel approved.xlsx --fixture-dir fixtures_mutated --apply  # time this
"""
import argparse, random, re, shutil, sys, time, uuid
from pathlib import Path

CONTENT_SELECTOR = "main, article, .content, #content, .page-content, .main-content, [role='main']"  # same as freshness
UA = {"User-Agent": "Mozilla/5.0"}


def normalize_url(url):  # same as freshness/scraper (slug must match fetch_html)
    url = url.strip().split("?")[0].split("#")[0].rstrip("/")
    if "://" in url:
        scheme_host, _, path = url.partition("://")
        i = path.find("/")
        if i == -1:
            return f"{scheme_host}://{path}"
        return f"{scheme_host}://{path[:i]}{path[i:].lower()}"
    return url.lower()


def slug(url):
    return re.sub(r"[^a-zA-Z0-9]+", "_", normalize_url(url)).strip("_")


def read_urls(excel):
    from openpyxl import load_workbook
    ws = load_workbook(excel, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c or "").strip().lower() for c in next(rows)]
    col = next((i for i, h in enumerate(hdr) if h in ("url", "urls", "page url", "link")), None)
    if col is None:
        sys.exit(f"No URL column in {hdr}")
    seen, out = set(), []
    for r in rows:
        u = str(r[col] or "").strip()
        if u.startswith("http") and normalize_url(u) not in seen:
            seen.add(normalize_url(u)); out.append(u)
    return out


def download(a):
    import requests
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    urls, failed = read_urls(a.excel), []
    for n, u in enumerate(urls, 1):
        for attempt in range(1, 4):
            try:
                r = requests.get(u, headers=UA, timeout=30)
                if r.status_code >= 500:
                    raise RuntimeError(f"status {r.status_code}")
                if r.status_code >= 400:
                    failed.append((u, f"status {r.status_code}")); break
                (out / f"{slug(u)}.html").write_text(r.text, encoding="utf-8")
                break
            except Exception as e:
                if attempt == 3:
                    failed.append((u, str(e)))
                else:
                    time.sleep(2 ** attempt)
        print(f"[{n}/{len(urls)}] {u}", flush=True)
        time.sleep(a.delay)
    print(f"saved {len(urls) - len(failed)}/{len(urls)} to {out}")
    for u, why in failed:
        print("FAILED", u, why)
    sys.exit(2 if failed else 0)


def mutate(a):
    from bs4 import BeautifulSoup
    src, dst = Path(a.src), Path(a.dst)
    if dst.exists():
        shutil.rmtree(dst)
    dst.mkdir(parents=True)
    files = sorted(src.glob("*.html"))
    if not files:
        sys.exit(f"no .html in {src}")
    random.seed(a.seed)
    chosen = set(random.sample(files, max(1, round(len(files) * a.percent / 100))))
    run = uuid.uuid4().hex[:8]
    changed = 0
    for f in files:
        html = f.read_text(encoding="utf-8")
        if f in chosen:
            soup = BeautifulSoup(html, "html.parser")
            box = soup.select_one(CONTENT_SELECTOR) or soup.body or soup
            p = soup.new_tag("p")
            p.string = f"Load test update {run}-{f.stem[-24:]}: this sentence was added to force a content change."
            box.append(p)
            html = str(soup); changed += 1
        (dst / f.name).write_text(html, encoding="utf-8")
    print(f"copied {len(files)} files, mutated {changed} ({a.percent}%) -> {dst} (run {run})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("download"); d.add_argument("--excel", required=True); d.add_argument("--out", required=True)
    d.add_argument("--delay", type=float, default=0.5); d.set_defaults(fn=download)
    m = sub.add_parser("mutate"); m.add_argument("--src", required=True); m.add_argument("--dst", required=True)
    m.add_argument("--percent", type=float, default=100); m.add_argument("--seed", type=int, default=1); m.set_defaults(fn=mutate)
    a = ap.parse_args(); a.fn(a)