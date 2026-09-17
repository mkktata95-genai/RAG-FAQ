"""
extract_images_for_url_list.py

Reads URLs from an Excel file, extracts each page's citation thumbnail
image (base_image_url) using the finalized extraction logic, and writes
the result into a new column immediately next to the URL column. Leaves
the cell blank if no image was found or the page failed to fetch.

Tries four patterns per page, in order (see CHANGELOG for how each was
found):
  1. JSON-LD "image" array (Feature Article Page template).
  2. <picture><source srcset> hero banner (id="herobannerblock*").
  3. CSS background-image hero banner (id="featurebanner*").
  4. imageblock content image (id="image*") — pages with no hero banner
     at all; anchored to the FIRST such block only.
og:image is deliberately NOT used — confirmed always generic (site logo
or teaser stock) across every page checked.

CONFIGURE these before running:
    INPUT_FILE     - path to your Excel file
    SHEET_NAME     - sheet name, or None for the active sheet
    URL_COLUMN     - the exact header text of the column containing URLs

Output: writes a new file "<original_name>_with_images.xlsx" next to the
input file. Does not modify the original file.

Usage:
    python extract_images_for_url_list.py

CHANGELOG
---------
v3: Scoped extraction to the banner container (id-anchored) instead of
    whole-page scan; added <picture><source> support; fixed a
    nested-paren regex bug on linear-gradient(rgba(...)) overlays.
v4: Added JSON-LD "image" array as primary source (Feature Article Page
    template) — fixed the majority of a 158/294 blank-result batch run.
v5: Added "imageblock" (id="image*") fallback for pages with no hero
    banner section at all (e.g. find-a-financial-adviser).
v6: Removed file-extension requirement on CSS url() matches — real
    extension-less URLs found in production (default banner assets that
    still render live despite no .jpg/.png suffix).
v7: Filter generic/fallback filenames out of JSON-LD "image" results
    (e.g. "feature-article-fallback-...jpg" on articles with no custom
    banner set) instead of trusting JSON-LD blindly; falls through to
    other sources or leaves the cell blank.
v8: Skip the imageblock fallback entirely on Feature Article Page
    templates (detected via JSON-LD "@type":"Article" or meta
    template=FeatureArticlePageType) — their id="image*" blocks are
    mid-body content illustrations, not a hero image, and were being
    wrongly picked up when JSON-LD had no real image.
v9: Generic-filename filter now applies to every source's candidate, not
    just JSON-LD. Confirmed real case: "couple-in-discussion-
    small-550x550.jpg" — a shared "Do you need financial support?" CTA
    embedded on multiple unrelated articles — was reused via the same
    id="featurebanner*" naming as a real hero banner, so it was returned
    as if it were the page's own image.
"""

import re
import json
import time
from urllib.parse import urljoin
from typing import Optional

import requests
import openpyxl

# ---------------------------------------------------------------------------
# CONFIGURE
# ---------------------------------------------------------------------------
INPUT_FILE = "Approved_URLS_140826_checked.xlsx"
SHEET_NAME = None          # e.g. "Sheet1", or None for the active sheet
URL_COLUMN = "URL"         # exact header text of the column with page URLs
NEW_COLUMN_HEADER = "base_image_url"
REQUEST_TIMEOUT = 15
DELAY_BETWEEN_REQUESTS = 0.5   # seconds, be polite to the site
# ---------------------------------------------------------------------------

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0 Safari/537.36"
}

_LDJSON_RE = re.compile(
    r'<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
    re.IGNORECASE | re.DOTALL,
)

# v9: checked against every source's candidate now, not just JSON-LD.
_GENERIC_IMAGE_PATTERNS = ["fallback", "rl_logo", "rl-logo", "couple-in-discussion"]


def _is_generic_image(path: str) -> bool:
    lower = path.lower()
    return any(p in lower for p in _GENERIC_IMAGE_PATTERNS)


def _extract_from_jsonld(html: str) -> Optional[str]:
    for match in _LDJSON_RE.finditer(html):
        raw = match.group(1).strip()
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            continue
        candidates = data if isinstance(data, list) else [data]
        for item in candidates:
            if not isinstance(item, dict):
                continue
            image = item.get("image")
            if not image:
                continue
            if isinstance(image, list) and image:
                chosen = None
                for img in image:
                    if isinstance(img, str) and "large" in img.lower():
                        chosen = img
                        break
                if chosen is None:
                    last = image[-1]
                    chosen = last if isinstance(last, str) else None
                if chosen and not _is_generic_image(chosen):
                    return chosen
                continue
            if isinstance(image, str) and not _is_generic_image(image):
                return image
    return None


_BANNER_ANCHOR_RE = re.compile(
    r'id=["\'](?:feature|hero)banner(?:block)?\d+["\']', re.IGNORECASE
)
_DECLARATION_RE = re.compile(r'background-image:\s*([^;]+);', re.IGNORECASE)
# no extension required — some pages use a genuinely extension-less URL
# for a default/placeholder banner asset that still renders live.
_URL_RE = re.compile(r'url\(\s*[\'"]?([^\'")]+)[\'"]?\s*\)', re.IGNORECASE)
_SOURCE_TAG_RE = re.compile(r'<source\b[^>]*>', re.IGNORECASE)
_SRCSET_RE = re.compile(
    r'srcset=["\']([^"\']+\.(?:jpg|jpeg|png|webp))["\']', re.IGNORECASE
)
_IMG_SRC_RE = re.compile(
    r'<img\b[^>]*\bsrc=["\']([^"\']+\.(?:jpg|jpeg|png|webp))["\']', re.IGNORECASE
)
SIZE_TIER_PREFERENCE = ["medium-1000x570", "small-800x800", "large-1200x700"]
_WINDOW_FALLBACK_SIZE = 8000
_IMAGEBLOCK_ANCHOR_RE = re.compile(r'id=["\']image\d+["\']', re.IGNORECASE)
_IMAGEBLOCK_WINDOW_FALLBACK = 3000


def _get_imageblock_window(html: str) -> Optional[str]:
    anchor = _IMAGEBLOCK_ANCHOR_RE.search(html)
    if not anchor:
        return None
    start = anchor.end()
    figure_end = html.find("</figure>", start)
    end = figure_end if figure_end != -1 else start + _IMAGEBLOCK_WINDOW_FALLBACK
    return html[start:end]


def _extract_from_imageblock(window: str) -> Optional[str]:
    picture_result = _extract_from_picture(window)
    if picture_result:
        return picture_result
    img_match = _IMG_SRC_RE.search(window)
    if img_match:
        return img_match.group(1)
    return None


def _get_banner_window(html: str) -> Optional[str]:
    anchor = _BANNER_ANCHOR_RE.search(html)
    if not anchor:
        return None
    start = anchor.end()
    next_container = html.find('<div class="container', start + 100)
    end = next_container if next_container != -1 else start + _WINDOW_FALLBACK_SIZE
    return html[start:end]


def _extract_from_picture(window: str) -> Optional[str]:
    candidates = []
    for tag in _SOURCE_TAG_RE.findall(window):
        srcset_match = _SRCSET_RE.search(tag)
        if not srcset_match:
            continue
        has_media = "media=" in tag.lower()
        candidates.append((srcset_match.group(1), has_media))
    if not candidates:
        return None
    for path, has_media in candidates:
        if not has_media:
            return path
    return candidates[0][0]


def _extract_from_css_banner(window: str) -> Optional[str]:
    matches = []
    seen = set()
    for decl_match in _DECLARATION_RE.finditer(window):
        declaration = decl_match.group(1)
        urls_in_declaration = _URL_RE.findall(declaration)
        if not urls_in_declaration:
            continue
        rel_path = urls_in_declaration[-1]
        if rel_path not in seen:
            seen.add(rel_path)
            matches.append(rel_path)
    if not matches:
        return None
    for tier in SIZE_TIER_PREFERENCE:
        for rel_path in matches:
            if tier in rel_path:
                return rel_path
    return matches[0]


# v8: skip imageblock fallback on Feature Article Page templates — their
# id="image*" blocks are mid-body content illustrations (e.g. a table
# screenshot), not a hero/citation image. Confirmed real case:
# "the-royal-london-guide-to-going-self-employed" has no hero image at
# all, and the old logic fell through to a random body table.png instead
# of correctly returning blank.
_ARTICLE_JSONLD_TYPE_RE = re.compile(r'"@type"\s*:\s*"Article"', re.IGNORECASE)
_ARTICLE_META_TEMPLATE_RE = re.compile(
    r'<meta\s+name=["\']template["\'][^>]*content=["\']FeatureArticlePageType["\']',
    re.IGNORECASE,
)


def _is_article_template(html: str) -> bool:
    return bool(
        _ARTICLE_JSONLD_TYPE_RE.search(html)
        or _ARTICLE_META_TEMPLATE_RE.search(html)
    )


def extract_base_image_url(html: str, page_url: str) -> Optional[str]:
    """
    Tries four patterns in order: JSON-LD image array, picture/source hero,
    CSS background-image hero, imageblock content image (no-hero pages,
    skipped entirely on Feature Article Page templates). v9: every
    candidate is checked against _is_generic_image before being accepted
    — a generic match falls through to the next source. Does NOT fall
    back to scanning the rest of the page.
    """
    jsonld_result = _extract_from_jsonld(html)
    if jsonld_result:
        return urljoin(page_url, jsonld_result)

    window = _get_banner_window(html)
    if window is not None:
        picture_result = _extract_from_picture(window)
        if picture_result and not _is_generic_image(picture_result):
            return urljoin(page_url, picture_result)
        css_result = _extract_from_css_banner(window)
        if css_result and not _is_generic_image(css_result):
            return urljoin(page_url, css_result)

    if _is_article_template(html):
        return None

    imageblock_window = _get_imageblock_window(html)
    if imageblock_window is not None:
        imageblock_result = _extract_from_imageblock(imageblock_window)
        if imageblock_result and not _is_generic_image(imageblock_result):
            return urljoin(page_url, imageblock_result)

    return None


def fetch_image_for_url(url: str) -> Optional[str]:
    try:
        resp = requests.get(url, headers=HEADERS, timeout=REQUEST_TIMEOUT)
        resp.raise_for_status()
    except Exception as e:
        print(f"    ERROR fetching {url}: {e}")
        return None
    return extract_base_image_url(resp.text, url)


def main():
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active

    # find URL column and header row
    header_row = 1
    headers = {}
    for cell in ws[header_row]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    if URL_COLUMN not in headers:
        print(f"Column '{URL_COLUMN}' not found. Available columns: {list(headers.keys())}")
        return

    url_col_idx = headers[URL_COLUMN]

    # insert the new column right after the URL column
    new_col_idx = url_col_idx + 1
    ws.insert_cols(new_col_idx)
    ws.cell(row=header_row, column=new_col_idx, value=NEW_COLUMN_HEADER)

    total_rows = ws.max_row - header_row
    found_count = 0

    for row_idx in range(header_row + 1, ws.max_row + 1):
        url_cell = ws.cell(row=row_idx, column=url_col_idx)
        url = url_cell.value
        if not url or not str(url).strip().startswith("http"):
            continue

        url = str(url).strip()
        print(f"[{row_idx - header_row}/{total_rows}] {url}")

        image_url = fetch_image_for_url(url)
        if image_url:
            found_count += 1
            print(f"    -> {image_url}")
        else:
            print(f"    -> (none found)")

        ws.cell(row=row_idx, column=new_col_idx, value=image_url or "")

        time.sleep(DELAY_BETWEEN_REQUESTS)

    output_file = INPUT_FILE.rsplit(".", 1)[0] + "_with_images.xlsx"
    wb.save(output_file)

    print("\n" + "=" * 70)
    print(f"Done. {found_count}/{total_rows} URLs had an image found.")
    print(f"Saved to: {output_file}")
    print("=" * 70)


if __name__ == "__main__":
    main()