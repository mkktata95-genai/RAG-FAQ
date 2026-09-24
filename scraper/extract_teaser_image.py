"""
extract_teaser_image.py

Reads URLs from an Excel file, extracts each page's official CMS-provided
"teaser_image" metadata field (Swifttype meta tag), and writes the result
into a new column immediately next to the URL column. Leaves the cell
blank if the tag is absent or the page failed to fetch.

Why this replaces the earlier hero-banner-scraping approach
-------------------------------------------------------------
B&M (Jodi) confirmed there's a single, CMS-maintained metadata field meant
for exactly this purpose (content-card thumbnails site-wide):

    <meta class="swifttype" name="teaser_image" data-type="enum"
          content="https://www.royallondon.com/globalassets/.../xxx-350x200.jpg" />

This is simpler and more reliable than reverse-engineering hero banners
from page HTML/CSS (which needed v3-v9 of fixes to handle template
differences, generic/reused images, and extension-less URLs). Confirmed
live on two pages that were problem cases for the old approach:
  - why-your-credit-rating-is-important (old logic needed a blocklist to
    reject a reused CTA image here) -> teaser_image correctly returns
    newspaper-and-glasses-350x200.jpg
  - find-a-financial-adviser (StandardPageType, old logic assumed "no
    hero image" and used a risky imageblock fallback) -> teaser_image is
    present: office-teaser-350x200.jpg

Not every page is guaranteed to have this tag populated (per B&M) — the
script leaves those blank so gaps can be measured and a fallback image
strategy sized accordingly.

CONFIGURE these before running:
    INPUT_FILE     - path to your Excel file
    SHEET_NAME     - sheet name, or None for the active sheet
    URL_COLUMN     - the exact header text of the column containing URLs

Output: writes a new file "<original_name>_with_teaser_images.xlsx" next
to the input file. Does not modify the original file.

Usage:
    python extract_teaser_image.py
"""

import re
import time
from typing import Optional

import requests
import openpyxl

# ---------------------------------------------------------------------------
# CONFIGURE
# ---------------------------------------------------------------------------
INPUT_FILE = "Approved_URLS_140826_checked.xlsx"
SHEET_NAME = None          # e.g. "Sheet1", or None for the active sheet
URL_COLUMN = "URL"         # exact header text of the column with page URLs
NEW_COLUMN_HEADER = "teaser_image_url"
REQUEST_TIMEOUT = 15
DELAY_BETWEEN_REQUESTS = 0.5   # seconds, be polite to the site
RUN_REGRESSION_TESTS = True    # set False to skip straight to the batch run
# ---------------------------------------------------------------------------

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0 Safari/537.36"
}

# Matches the swifttype teaser_image meta tag regardless of attribute
# order (content before/after name, single or double quotes, self-closing
# or not). Anchored on name="teaser_image" so it doesn't accidentally
# match teaser_title / teaser_text.
_TEASER_META_RE = re.compile(
    r'<meta\b(?=[^>]*\bname=["\']teaser_image["\'])(?=[^>]*\bcontent=["\']([^"\']+)["\'])[^>]*>',
    re.IGNORECASE,
)


def extract_teaser_image(html: str) -> Optional[str]:
    match = _TEASER_META_RE.search(html)
    if match:
        return match.group(1).strip()
    return None


def fetch_teaser_image(url: str) -> Optional[str]:
    try:
        resp = requests.get(url, headers=HEADERS, timeout=REQUEST_TIMEOUT)
        resp.raise_for_status()
    except Exception as e:
        print(f"    ERROR fetching {url}: {e}")
        return None
    return extract_teaser_image(resp.text)


def main():
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active

    header_row = 1
    headers = {}
    for cell in ws[header_row]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    if URL_COLUMN not in headers:
        print(f"Column '{URL_COLUMN}' not found. Available columns: {list(headers.keys())}")
        return

    url_col_idx = headers[URL_COLUMN]
    new_col_idx = url_col_idx + 1
    ws.insert_cols(new_col_idx)
    ws.cell(row=header_row, column=new_col_idx, value=NEW_COLUMN_HEADER)

    total_rows = ws.max_row - header_row
    found_count = 0

    for row_idx in range(header_row + 1, ws.max_row + 1):
        url_cell = ws.cell(row=row_idx, column=url_col_idx)
        url = url_cell.value
        if not url or not str(url).strip().startswith("http"):
            continue

        url = str(url).strip()
        print(f"[{row_idx - header_row}/{total_rows}] {url}")

        teaser_url = fetch_teaser_image(url)
        if teaser_url:
            found_count += 1
            print(f"    -> {teaser_url}")
        else:
            print(f"    -> (none found)")

        ws.cell(row=row_idx, column=new_col_idx, value=teaser_url or "")

        time.sleep(DELAY_BETWEEN_REQUESTS)

    output_file = INPUT_FILE.rsplit(".", 1)[0] + "_with_teaser_images.xlsx"
    wb.save(output_file)

    print("\n" + "=" * 70)
    print(f"Done. {found_count}/{total_rows} URLs had a teaser_image found.")
    print(f"Coverage: {found_count}/{total_rows} = {found_count/total_rows*100:.1f}%")
    print(f"Saved to: {output_file}")
    print("=" * 70)


def run_regression_tests() -> bool:
    """Returns True if all tests pass. Does NOT run the batch job itself —
    call main() separately (see bottom of file)."""
    # Regression tests against real tag formats confirmed live on
    # royallondon.com (see module docstring).
    tests = []

    # 1) Real confirmed case: how-to-boost-your-pension-pot (from B&M's screenshot)
    html1 = '''
    <meta class="swifttype" name="st-description" data-type="string" content="Discover how to increase your pension contributions to help boost your pension pot and fund your retirement. Learn more on the Royal London website." />
    <meta class="swifttype" name="teaser_image" data-type="enum" content="https://www.royallondon.com/globalassets/all-site-images/brand-photography/teaser-images/tea-teaser-350x200.jpg" />
    <meta class="swifttype" name="teaser_title" data-type="string" content="How to boost your pension pot" />
    <meta class="swifttype" name="teaser_text" data-type="string" content="Discover how to increase your pension contributions to help boost your pension pot and fund your retirement." />
    '''
    tests.append(("pension pot page (present)", html1,
                  "https://www.royallondon.com/globalassets/all-site-images/brand-photography/teaser-images/tea-teaser-350x200.jpg"))

    # 2) Real confirmed case: why-your-credit-rating-is-important (old logic's
    #    generic-CTA-image bug case) - teaser_image is correct here, unaffected
    #    by the "Do you need financial support?" CTA image reuse.
    html2 = '''
    <meta class="swifttype" name="teaser_image" data-type="enum" content="https://www.royallondon.com/globalassets/all-site-images/brand-photography/teaser-images/newspaper-and-glasses-350x200.jpg" />
    '''
    tests.append(("credit-rating page (present, correct despite CTA reuse elsewhere)", html2,
                  "https://www.royallondon.com/globalassets/all-site-images/brand-photography/teaser-images/newspaper-and-glasses-350x200.jpg"))

    # 3) Real confirmed case: find-a-financial-adviser (StandardPageType -
    #    old logic assumed no hero image existed here)
    html3 = '''
    <meta class="swifttype" name="teaser_image" data-type="enum" content="https://www.royallondon.com/globalassets/all-site-images/brand-photography/teaser-images/office-teaser-350x200.jpg" />
    '''
    tests.append(("find-a-financial-adviser (present)", html3,
                  "https://www.royallondon.com/globalassets/all-site-images/brand-photography/teaser-images/office-teaser-350x200.jpg"))

    # 4) Tag absent entirely -> must return None, not raise
    html4 = '''
    <meta class="swifttype" name="st-description" data-type="string" content="Some other page with no teaser image set." />
    '''
    tests.append(("no teaser_image tag (absent)", html4, None))

    # 5) Attribute order reversed (content before name) - some CMS renders
    #    can vary attribute order; must still match.
    html5 = '''
    <meta content="https://www.royallondon.com/globalassets/xyz-teaser-350x200.jpg" data-type="enum" name="teaser_image" class="swifttype" />
    '''
    tests.append(("reversed attribute order", html5,
                  "https://www.royallondon.com/globalassets/xyz-teaser-350x200.jpg"))

    # 6) Must not accidentally match teaser_title / teaser_text
    html6 = '''
    <meta class="swifttype" name="teaser_title" data-type="string" content="Should not match" />
    <meta class="swifttype" name="teaser_text" data-type="string" content="Should not match either" />
    '''
    tests.append(("only teaser_title/teaser_text present, no teaser_image", html6, None))

    passed = 0
    for name, html, expected in tests:
        result = extract_teaser_image(html)
        ok = result == expected
        status = "PASS" if ok else "FAIL"
        if ok:
            passed += 1
        print(f"[{status}] {name}\n    expected: {expected}\n    got:      {result}\n")

    print(f"{passed}/{len(tests)} tests passed\n")
    return passed == len(tests)


if __name__ == "__main__":
    if RUN_REGRESSION_TESTS:
        if not run_regression_tests():
            print("Regression tests FAILED — fix before running the batch job. Aborting.")
            raise SystemExit(1)
        print("All regression tests passed. Starting batch run against the URL list...\n")

    main()