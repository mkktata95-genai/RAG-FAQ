from azure.storage.blob import BlobServiceClient
from openpyxl import load_workbook
import io

c = BlobServiceClient.from_connection_string("UseDevelopmentStorage=true").get_container_client("output")
name = sorted(b.name for b in c.list_blobs(name_starts_with="freshness-reports/"))[-1]
print(name)
wb = load_workbook(io.BytesIO(c.download_blob(name).readall()), read_only=True)
for ws in wb.worksheets:
    print("--", ws.title)
    for row in ws.iter_rows(values_only=True):
        cells = [str(x) for x in row if x is not None]
        if any(x in ("changed", "new") or x.startswith("removed_") or x == "scrape_failed" for x in cells):
            print(cells)