"""Templated-questions builder: Excel -> validated rows -> NEW timestamped Search index.

Same model as the indexer: a build never touches the live index; promote() points the
alias at the new index (manual, reversible). Index schema is the agent's `scache`
schema (id, question, question_vector, response_json, created_at) plus `source_urls`.

The agent parses response_json with LLMResponse (extra="forbid") and treats a parse
failure as a SILENT cache miss, so validate() mirrors those rules and a single bad
row refuses the whole build (nothing is created).
"""
import hashlib
import json
import logging
import os
import re
import tempfile
import threading
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

log = logging.getLogger("templated_job")

# ── settings (all overridable by app settings) ────────────────
INPUT_PREFIX = os.environ.get("TEMPLATED_INPUT_PREFIX", "templated-questions/")
ALIAS_NAME = os.environ.get("TEMPLATED_ALIAS_NAME", "rlg-faq-templated-live")
KEEP_INDEXES = int(os.environ.get("TEMPLATED_KEEP_COUNT", "3"))
MIN_PROMOTE_RATIO = float(os.environ.get("TEMPLATED_MIN_PROMOTE_RATIO", "0.5"))
MIN_DOCS_FOR_GATE = 20
INDEX_RE = re.compile(r"^rlg-faq-templated-\d{8}-\d{6}$")
# Agent rules (config.yaml prompt.answer + src/core/constant.py). Keep in sync.
ANSWER_HARD_MAX = int(os.environ.get("TEMPLATED_ANSWER_MAX_CHARS", "250"))  # char_hard_limit 200 + char_buffer 50
ANSWER_SOFT = int(os.environ.get("TEMPLATED_ANSWER_SOFT_CHARS", "180"))     # char_soft_limit (warn only)
CITATION_MIN, CITATION_MAX = int(os.environ.get("TEMPLATED_CITATION_MIN", "0")), int(os.environ.get("TEMPLATED_CITATION_MAX", "3"))
READMORE_MIN, READMORE_MAX = int(os.environ.get("TEMPLATED_READMORE_MIN", "0")), int(os.environ.get("TEMPLATED_READMORE_MAX", "3"))
INTENTS = {"blocked", "direct", "grounded"}
SCENARIOS = {"not_applicable", "standard", "advice_seeking", "distress", "escalation_and_handoff", "account_action"}
ALLOWED_HOSTS = tuple(h.strip().lower() for h in os.environ.get("TEMPLATED_ALLOWED_HOSTS", "royallondon.com").split(",") if h.strip())
ID_RE = re.compile(r"^[A-Za-z0-9\-][A-Za-z0-9_=\-]*$")  # same as scache ingest
REQUIRED_ENV = ("AZURE_SEARCH_ENDPOINT", "AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_EMBEDDING_DEPLOYMENT",
                "AZURE_OPENAI_EMBEDDING_DIMENSIONS")

_JOB_LOCK = threading.Lock()


# ── Excel -> rows ─────────────────────────────────────────────
def read_excel(path: str) -> list[dict]:
    """First sheet; header row = column names (case/space-insensitive). Blank rows skipped."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [re.sub(r"\s+", "_", str(h).strip().lower()) if h is not None else "" for h in next(it, ())]
    rows = []
    for n, vals in enumerate(it, start=2):  # n = Excel row number
        row = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, vals) if h}
        if any(row.values()):
            row["_row"] = n
            rows.append(row)
    wb.close()
    return rows


# ── validation ────────────────────────────────────────────────
def norm_question(q: str) -> str:
    return re.sub(r"\s+", " ", q).strip().rstrip("?").strip().casefold()


def make_id(q: str) -> str:
    return "q-" + hashlib.sha1(norm_question(q).encode("utf-8")).hexdigest()[:10]


def _url_problem(url: str) -> str | None:
    p = urlparse(url)
    if p.scheme != "https" or not p.netloc:
        return "must be an absolute https URL"
    host = (p.hostname or "").lower()
    if ALLOWED_HOSTS and not any(host == h or host.endswith("." + h) for h in ALLOWED_HOSTS):
        return f"host '{host}' is not allowed (only {', '.join(ALLOWED_HOSTS)})"
    return None


def _pairs(row: dict, a: str, b: str, prefix: str, kind: str, lo: int, hi: int, errs: list):
    out = []
    for i in range(1, 10):
        x, y = row.get(f"{prefix}_{i}_{a}", ""), row.get(f"{prefix}_{i}_{b}", "")
        if not x and not y:
            continue
        if i > hi:
            errs.append(f"{kind} {i}: at most {hi} allowed")
            continue
        if not x or not y:
            errs.append(f"{kind} {i}: both {a} and {b} are required")
            continue
        bad = _url_problem(y)
        if bad:
            errs.append(f"{kind} {i} url {bad}")
            continue
        out.append({a: x, b: y})
    if len(out) < lo:
        errs.append(f"at least {lo} {kind} required")
    return out


def _bool(v: str, default=False):
    if v == "":
        return default
    if v.lower() in ("true", "yes", "y", "1"):
        return True
    if v.lower() in ("false", "no", "n", "0"):
        return False
    raise ValueError(v)


def validate(rows: list[dict]) -> tuple[list[dict], list[str], list[str]]:
    """-> (entries, errors, warnings). Any error means the build must be refused."""
    entries, errors, warnings, seen_q, seen_id = [], [], [], {}, {}
    for row in rows:
        n, errs = row["_row"], []
        if row.get("active", "y").lower() in ("n", "no", "false", "0"):
            continue  # retired row: simply not indexed
        q, ans = row.get("question", ""), row.get("answer", "")
        if not q:
            errs.append("question is empty")
        if not ans:
            errs.append("answer is empty")
        elif len(ans) > ANSWER_HARD_MAX:
            errs.append(f"answer is {len(ans)} chars, over the {ANSWER_HARD_MAX} limit")
        elif len(ans) > ANSWER_SOFT:
            warnings.append(f"row {n}: answer is {len(ans)} chars (soft limit {ANSWER_SOFT})")
        intent, scenario = row.get("intent", "") or "grounded", row.get("scenario", "") or "standard"
        if intent not in INTENTS:
            errs.append(f"intent '{intent}' not in {sorted(INTENTS)}")
        if scenario not in SCENARIOS:
            errs.append(f"scenario '{scenario}' not in {sorted(SCENARIOS)}")
        try:
            boundary = _bool(row.get("advice_boundary", ""))
        except ValueError:
            errs.append("advice_boundary must be true/false")
            boundary = False
        cites = _pairs(row, "title", "url", "citation", "citation", CITATION_MIN, CITATION_MAX, errs)
        ctas = _pairs(row, "label", "url", "cta", "cta", READMORE_MIN, READMORE_MAX, errs)
        eid = row.get("id", "") or (make_id(q) if q else "")
        if row.get("id") and not ID_RE.match(eid):
            errs.append(f"id '{eid}' has invalid characters")
        if q:
            k = norm_question(q)
            if k in seen_q:
                errs.append(f"duplicate of row {seen_q[k]}")
            seen_q.setdefault(k, n)
        if eid:
            if eid in seen_id:
                errs.append(f"id '{eid}' already used in row {seen_id[eid]}")
            seen_id.setdefault(eid, n)
        if errs:
            errors += [f"row {n}: {e}" for e in errs]
            continue
        entries.append({
            "id": eid, "question": q,
            # key order == LLMResponse field order == model_dump_json order
            "response_json": json.dumps({"intent": intent, "scenario": scenario, "answer": ans,
                                         "advice_boundary": boundary, "citations": cites, "read_more": ctas},
                                        ensure_ascii=False, separators=(",", ":")),
            "source_urls": sorted({c["url"] for c in cites} | {c["url"] for c in ctas}),
        })
    if not entries and not errors:
        errors.append("no active rows found")
    return entries, errors, warnings


# ── embeddings ────────────────────────────────────────────────
def azure_embed(texts: list[str]) -> list[list[float]]:
    from azure.identity import DefaultAzureCredential, get_bearer_token_provider
    from openai import AzureOpenAI
    client = AzureOpenAI(
        azure_endpoint=os.environ["AZURE_OPENAI_ENDPOINT"],
        api_version=os.environ.get("AZURE_OPENAI_API_VERSION", "2024-10-21"),
        azure_ad_token_provider=get_bearer_token_provider(DefaultAzureCredential(), "https://cognitiveservices.azure.com/.default"),
    )
    dims = int(os.environ["AZURE_OPENAI_EMBEDDING_DIMENSIONS"])
    out = []
    for i in range(0, len(texts), 16):
        r = client.embeddings.create(model=os.environ["AZURE_OPENAI_EMBEDDING_DEPLOYMENT"], input=texts[i:i + 16], dimensions=dims)
        out += [d.embedding for d in sorted(r.data, key=lambda d: d.index)]
    return out


# ── Azure Search operations (faked in tests) ──────────────────
class AzureSearchOps:
    def __init__(self):
        from azure.identity import DefaultAzureCredential
        from azure.search.documents.indexes import SearchIndexClient
        self._endpoint = os.environ.get("AZURE_SEARCH_ENDPOINT", "").rstrip("/")
        if not self._endpoint:
            raise RuntimeError("AZURE_SEARCH_ENDPOINT is not set")
        self._cred = DefaultAzureCredential()
        self._ic = SearchIndexClient(self._endpoint, self._cred)

    def list_indexes(self) -> list[str]:
        return list(self._ic.list_index_names())

    def count(self, index: str) -> int:
        from azure.search.documents import SearchClient
        return SearchClient(self._endpoint, index, self._cred).get_document_count()

    def get_alias(self) -> str | None:
        from azure.core.exceptions import ResourceNotFoundError
        try:
            idx = self._ic.get_alias(ALIAS_NAME).indexes
        except ResourceNotFoundError:
            return None
        return idx[0] if idx else None

    def set_alias(self, index: str) -> None:
        from azure.search.documents.indexes.models import SearchAlias
        self._ic.create_or_update_alias(SearchAlias(name=ALIAS_NAME, indexes=[index]))

    def delete(self, index: str) -> None:
        self._ic.delete_index(index)

    def create(self, index: str, dims: int) -> None:
        """Same as the agent's scache index.json (cosine is required: agent computes 2 - 1/score) + source_urls."""
        from azure.search.documents.indexes.models import (
            HnswAlgorithmConfiguration, HnswParameters, SearchField, SearchFieldDataType, SearchIndex,
            VectorSearch, VectorSearchAlgorithmMetric, VectorSearchProfile)
        T = SearchFieldDataType
        fields = [
            SearchField(name="id", type=T.String, key=True, filterable=True),
            SearchField(name="question", type=T.String, searchable=True, retrievable=True),
            SearchField(name="question_vector", type=T.Collection(T.Single), searchable=True, retrievable=False,
                        stored=False, vector_search_dimensions=dims, vector_search_profile_name="cache-cosine-profile"),
            SearchField(name="response_json", type=T.String, searchable=False, retrievable=True),
            SearchField(name="source_urls", type=T.Collection(T.String), filterable=True, retrievable=True),
            SearchField(name="created_at", type=T.DateTimeOffset, filterable=True, sortable=True, retrievable=True),
        ]
        vs = VectorSearch(
            algorithms=[HnswAlgorithmConfiguration(name="cache-hnsw", parameters=HnswParameters(
                metric=VectorSearchAlgorithmMetric.COSINE, m=4, ef_construction=400, ef_search=500))],
            profiles=[VectorSearchProfile(name="cache-cosine-profile", algorithm_configuration_name="cache-hnsw")])
        self._ic.create_index(SearchIndex(name=index, fields=fields, vector_search=vs))

    def upload(self, index: str, docs: list[dict]) -> None:
        from azure.search.documents import SearchClient
        c = SearchClient(self._endpoint, index, self._cred)
        for i in range(0, len(docs), 100):
            res = c.upload_documents(docs[i:i + 100])
            bad = [r.key for r in res if not r.succeeded]
            if bad:
                raise RuntimeError(f"upload failed for ids {bad[:5]}")


# ── build ─────────────────────────────────────────────────────
def new_index_name(now: datetime | None = None) -> str:
    return "rlg-faq-templated-" + (now or datetime.now(timezone.utc)).strftime("%Y%m%d-%H%M%S")


def is_excel(name: str) -> bool:
    base = name.rsplit("/", 1)[-1]
    return base.lower().endswith(".xlsx") and not base.startswith(("~$", "."))  # skip Excel lock/temp files


def latest_excel(container) -> str:
    blobs = [b for b in container.list_blobs(name_starts_with=INPUT_PREFIX) if is_excel(b.name)]
    if not blobs:
        raise FileNotFoundError(f"No .xlsx found under input/{INPUT_PREFIX}")
    return max(blobs, key=lambda b: b.last_modified).name


def parse_build_message(body: str) -> str | None:
    try:
        name = json.loads(body).get("file")
        return (INPUT_PREFIX + name) if name else None
    except (ValueError, AttributeError):
        return None


def parse_promote_message(body: str) -> tuple[str | None, bool]:
    try:
        d = json.loads(body)
        if isinstance(d, dict):
            return d.get("index") or None, bool(d.get("force"))
    except ValueError:
        pass
    return None, "force" in body.lower()


def build_index(container_in, blob_name: str, embed=azure_embed, ops_factory=AzureSearchOps) -> dict:
    """Download Excel -> validate (no Azure calls yet) -> embed -> NEW index. Raises on any failure."""
    with _JOB_LOCK, tempfile.TemporaryDirectory(prefix="templated_") as tmp:
        src = Path(tmp) / Path(blob_name).name
        src.write_bytes(container_in.download_blob(blob_name).readall())
        entries, errors, warnings = validate(read_excel(str(src)))
        for w in warnings:
            log.warning("TEMPLATED_WARN %s", w)
        if errors:
            for e in errors:
                log.error("TEMPLATED_REJECTED %s: %s", blob_name, e)
            raise ValueError(f"{blob_name} rejected, {len(errors)} problem(s); first: {errors[0]}")

        index_name = new_index_name()
        dims = int(os.environ.get("AZURE_OPENAI_EMBEDDING_DIMENSIONS", "1536"))
        ops = ops_factory()
        try:
            vecs = embed([e["question"] for e in entries])
            if len(vecs) != len(entries) or any(len(v) != dims for v in vecs):
                raise RuntimeError(f"embedding result mismatch (expected {len(entries)} x {dims})")
            now = datetime.now(timezone.utc).isoformat()
            docs = [{"id": e["id"], "question": e["question"], "question_vector": v, "response_json": e["response_json"],
                     "source_urls": e["source_urls"], "created_at": now} for e, v in zip(entries, vecs)]
            ops.create(index_name, dims)
            ops.upload(index_name, docs)
        except Exception as ex:
            try:  # drop the half-built index
                ops.delete(index_name)
            except Exception:  # noqa: BLE001 - best effort
                pass
            raise RuntimeError(f"templated build failed ({index_name}): {ex}") from ex
        log.info("TEMPLATED_BUILT index=%s source=%s rows=%s (not live: run promote)", index_name, blob_name, len(entries))
        return {"index_name": index_name, "rows": len(entries), "warnings": warnings}


# ── promote (same rules as the indexer) ───────────────────────
def promote(ops, index: str | None = None, force: bool = False) -> dict:
    with _JOB_LOCK:
        built = sorted(n for n in ops.list_indexes() if INDEX_RE.match(n))
        target = index or (built[-1] if built else None)
        if not target or target not in built:
            log.error("PROMOTE_REFUSED index=%s not found among %s", target, built[-5:])
            return {"promoted": False, "reason": "index not found"}
        current = ops.get_alias()
        if current == target:
            return {"promoted": False, "reason": "already live", "index": target}
        if current and not force:
            n_new, n_cur = ops.count(target), ops.count(current)
            if n_cur >= MIN_DOCS_FOR_GATE and n_new < MIN_PROMOTE_RATIO * n_cur:
                log.error("PROMOTE_REFUSED %s has %s docs vs live %s has %s; send force to override", target, n_new, current, n_cur)
                return {"promoted": False, "reason": "shrink gate"}
        ops.set_alias(target)
        log.warning("TEMPLATED_PROMOTED alias=%s from=%s to=%s", ALIAS_NAME, current, target)
        keep = set(built[-KEEP_INDEXES:]) | {target} | ({current} if current else set())
        for name in built:
            if name not in keep:
                try:
                    ops.delete(name)
                except Exception as e:  # noqa: BLE001
                    log.warning("could not delete old index %s: %s", name, e)
        return {"promoted": True, "from": current, "index": target}
